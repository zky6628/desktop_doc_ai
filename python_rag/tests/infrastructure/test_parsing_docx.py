# -*- coding: utf-8 -*-
"""DOCX 解析器测试：顺序保真、表格证据、损坏/加密拒绝与确定性"""
import pytest
from docx import Document

from app.domain.errors import (
    DocumentCorruptedError,
    EmptyContentError,
    EncryptedDocumentError,
)
from app.domain.parsing import BlockType
from app.infrastructure.parsing.docx_parser import DocxParser

VERSION_ID = "01900000-0000-7000-8000-000000000001"


def _write_document(path) -> str:
    """生成含标题/段落/列表/表格交错内容的 DOCX 样例"""
    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("正文段落")
    document.add_paragraph("要点一", style="List Bullet")
    document.add_paragraph("要点二", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "表头甲"
    table.cell(0, 1).text = "表头乙"
    table.cell(1, 0).text = "数据甲"
    table.cell(1, 1).text = "数据乙"
    document.add_heading("第一节", level=2)
    document.add_paragraph("表格之后的段落")
    path = path.with_suffix(".docx")
    document.save(str(path))
    return str(path)


class TestBodyOrder:
    """主体元素顺序与块结构"""

    def test_paragraph_heading_list_table_order_preserved(self, tmp_path):
        parser = DocxParser()
        document = parser.parse(_write_document(tmp_path / "sample"), VERSION_ID)
        assert [block.block_type for block in document.blocks] == [
            BlockType.HEADING,
            BlockType.PARAGRAPH,
            BlockType.LIST,
            BlockType.TABLE,
            BlockType.HEADING,
            BlockType.PARAGRAPH,
        ]
        assert [block.ordinal for block in document.blocks] == [0, 1, 2, 3, 4, 5]
        assert document.blocks[0].text == "第一章"
        assert document.blocks[2].text == "要点一\n要点二"
        assert document.blocks[5].text == "表格之后的段落"
        assert all(block.page_no is None for block in document.blocks)

    def test_section_path_scopes_blocks(self, tmp_path):
        parser = DocxParser()
        document = parser.parse(_write_document(tmp_path / "sample"), VERSION_ID)
        assert document.blocks[1].section_path == "第一章"
        assert document.blocks[2].section_path == "第一章"
        assert document.blocks[3].section_path == "第一章"
        assert document.blocks[4].section_path == "第一章 > 第一节"
        assert document.blocks[5].section_path == "第一章 > 第一节"


class TestTableEvidence:
    """表格块的结构证据字段"""

    def test_table_evidence_fields(self, tmp_path):
        parser = DocxParser()
        document = parser.parse(_write_document(tmp_path / "sample"), VERSION_ID)
        table_block = document.blocks[3]
        evidence = table_block.table
        assert evidence is not None
        assert evidence.raw_html == (
            "<table><tr><td>表头甲</td><td>表头乙</td></tr>"
            "<tr><td>数据甲</td><td>数据乙</td></tr></table>"
        )
        assert evidence.raw_markdown == (
            "| 表头甲 | 表头乙 |\n| --- | --- |\n| 数据甲 | 数据乙 |"
        )
        assert evidence.structure_json == '{"rows":[["表头甲","表头乙"],["数据甲","数据乙"]]}'
        assert evidence.searchable_text == "表头甲 | 表头乙\n数据甲 | 数据乙"
        assert evidence.serialization_model == "local"
        assert evidence.serialization_version == "1"

    def test_pipe_in_cell_escaped(self, tmp_path):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "甲|乙"
        table.cell(0, 1).text = "丙"
        path = tmp_path / "pipe.docx"
        document.save(str(path))
        parsed = DocxParser().parse(str(path), VERSION_ID)
        evidence = parsed.blocks[0].table
        assert evidence is not None
        assert "甲\\|乙" in evidence.raw_markdown
        assert "| 丙 |" in evidence.raw_markdown

    def test_all_empty_table_skipped(self, tmp_path):
        document = Document()
        document.add_paragraph("唯一段落")
        document.add_table(rows=2, cols=2)
        path = tmp_path / "emptytable.docx"
        document.save(str(path))
        parsed = DocxParser().parse(str(path), VERSION_ID)
        assert [block.block_type for block in parsed.blocks] == [BlockType.PARAGRAPH]


class TestRejection:
    """损坏/加密/空内容的明确拒绝"""

    def test_ole_container_rejected_as_encrypted(self, tmp_path):
        path = tmp_path / "enc.docx"
        path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
        with pytest.raises(EncryptedDocumentError):
            DocxParser().parse(str(path), VERSION_ID)

    def test_garbage_bytes_rejected_as_corrupted(self, tmp_path):
        path = tmp_path / "bad.docx"
        path.write_bytes(b"this is not a zip archive at all")
        with pytest.raises(DocumentCorruptedError):
            DocxParser().parse(str(path), VERSION_ID)

    def test_zip_without_word_directory_rejected(self, tmp_path):
        import zipfile

        path = tmp_path / "fake.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("unrelated.txt", "content")
        with pytest.raises(DocumentCorruptedError):
            DocxParser().parse(str(path), VERSION_ID)

    def test_only_empty_paragraphs_rejected_as_empty(self, tmp_path):
        document = Document()
        document.add_paragraph("   ")
        document.add_paragraph("")
        path = tmp_path / "blank.docx"
        document.save(str(path))
        with pytest.raises(EmptyContentError):
            DocxParser().parse(str(path), VERSION_ID)

    def test_error_messages_contain_no_document_text(self, tmp_path):
        path = tmp_path / "bad.docx"
        path.write_bytes(b"garbage")
        try:
            DocxParser().parse(str(path), VERSION_ID)
        except DocumentCorruptedError as exc:
            message = str(exc)
            assert "garbage" not in message
            assert str(path) not in message


class TestDeterminism:
    """确定性合同：同文件重复解析输出逐字段一致"""

    def test_repeated_parse_identical(self, tmp_path):
        parser = DocxParser()
        path = _write_document(tmp_path / "sample")
        first = parser.parse(path, VERSION_ID)
        second = parser.parse(path, VERSION_ID)
        assert first == second
        assert first.structure_sha256 == second.structure_sha256
